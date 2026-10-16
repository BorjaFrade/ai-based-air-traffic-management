#!/usr/bin/env python3
"""
Download/compile ENAIRE AIP ENR 3.2 ATS corridors for one study area.

Default:
    python3 download_enaire_corridors.py --area LECMSAN

Default source:
    Current ENAIRE AIP ENR 3.2.

Historical studies:
    The default URL is NOT historical. For January 2019, use a historical
    ENR 3.2 HTML source if available:

        --source-url <historical_ENR_3_2_URL>

    or:

        --source-html /path/to/historical_ENR_3_2.html
"""

from __future__ import annotations

import argparse
import io
import json
import math
import re
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from pyproj import Geod

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


SCRIPT_DIR = Path(__file__).resolve().parent

if SCRIPT_DIR.name == "acquisition" and SCRIPT_DIR.parent.name == "src":
    PROJECT_ROOT = SCRIPT_DIR.parent.parent
elif SCRIPT_DIR.name == "acquisition":
    PROJECT_ROOT = SCRIPT_DIR.parent
else:
    PROJECT_ROOT = SCRIPT_DIR.parent

BBOX_PRESETS_FILE = SCRIPT_DIR / "bbox_presets.json"
DEFAULT_OUTPUT_FILE = SCRIPT_DIR / "corridors.json"
DEFAULT_WAYPOINTS_FILE = SCRIPT_DIR / "waypoints.json"

DEFAULT_SOURCE_URL = (
    "https://aip.enaire.es/AIP/contenido_AIP/ENR/"
    "LE_ENR_3_2_en.html"
)

DEFAULT_AREA = "LECMSAN"

GEOD = Geod(ellps="WGS84")

if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from area_presets import (  # noqa: E402
    load_area_presets,
    point_is_inside_area,
    resolve_area_preset,
)


ROUTE_RE = re.compile(r"^[A-Z]{1,2}[0-9]{1,4}[A-Z]?$")
COORD_RE = re.compile(
    r"(?P<lat>\d{6}(?:\.\d+)?[NS])\s+"
    r"(?P<lon>\d{7}(?:\.\d+)?[EW])"
)

SKIP_LABELS = {"▲", "∆", "△", "▼", "---"}
POINT_CODE_RE = re.compile(r"^[A-Z]{2,5}[0-9]?$")
AIRAC_VALIDITY_2019 = {
    "1813": ("2018-12-06", "2019-01-02"),
    "1901": ("2019-01-03", "2019-01-30"),
    "1902": ("2019-01-31", "2019-02-27"),
    "1903": ("2019-02-28", "2019-03-27"),
}


def parse_lat(value: str) -> float:
    value = value.strip().upper()
    h = value[-1]
    s = value[:-1]
    deg = int(s[:2])
    minute = int(s[2:4])
    sec = float(s[4:])
    out = deg + minute / 60.0 + sec / 3600.0
    return -out if h == "S" else out


def parse_lon(value: str) -> float:
    value = value.strip().upper()
    h = value[-1]
    s = value[:-1]
    deg = int(s[:3])
    minute = int(s[3:5])
    sec = float(s[5:])
    out = deg + minute / 60.0 + sec / 3600.0
    return -out if h == "W" else out


def compact_name(label: str) -> str:
    label = " ".join(str(label).split()).strip()

    codes = re.findall(
        r"\(([A-Z0-9]{2,5})\)",
        label,
    )

    if codes:
        return codes[-1]

    if not label:
        return ""

    return re.sub(
        r"[^A-Z0-9]",
        "",
        label.split()[0].upper(),
    )


def fetch_html(
    url: str,
    timeout_s: float,
) -> str:
    response = requests.get(
        url,
        headers={
            "User-Agent": (
                "Mozilla/5.0 "
                "(OpenSky academic ATS corridor downloader)"
            )
        },
        timeout=timeout_s,
    )
    response.raise_for_status()
    response.encoding = (
        response.apparent_encoding
        or response.encoding
        or "utf-8"
    )
    return response.text


def visible_lines(
    html: str,
) -> list[str]:
    soup = BeautifulSoup(
        html,
        "html.parser",
    )

    lines = []

    for raw in soup.get_text(
        "\n",
        strip=True,
    ).splitlines():
        line = " ".join(
            raw.split()
        ).strip()

        if line:
            lines.append(line)

    return lines


def previous_label(
    lines: list[str],
    index: int,
) -> str:
    for j in range(
        index - 1,
        max(-1, index - 8),
        -1,
    ):
        candidate = lines[j].strip()

        if not candidate:
            continue
        if candidate in SKIP_LABELS:
            continue
        if ROUTE_RE.fullmatch(candidate):
            continue
        if candidate.startswith("RNAV"):
            continue
        if candidate.startswith("FL"):
            continue

        return candidate

    return ""


def parse_enr32_routes(
    html: str,
) -> dict[str, list[dict]]:
    """
    Parse route blocks from the visible ENR 3.2 HTML text.

    A block is accepted only if at least two geographic points are found,
    which filters route-like tokens appearing inside explanatory notes.
    """
    lines = visible_lines(html)

    route_starts = [
        (index, line)
        for index, line in enumerate(lines)
        if ROUTE_RE.fullmatch(line)
    ]

    routes = {}

    for route_idx, (start, route_name) in enumerate(route_starts):
        end = (
            route_starts[route_idx + 1][0]
            if route_idx + 1 < len(route_starts)
            else len(lines)
        )

        points = []

        for index in range(start + 1, end):
            match = COORD_RE.search(lines[index])

            if not match:
                continue

            label = previous_label(
                lines,
                index,
            )

            if not label:
                continue

            try:
                point = {
                    "name": compact_name(label),
                    "label": label,
                    "lat": parse_lat(match.group("lat")),
                    "lon": parse_lon(match.group("lon")),
                }
            except Exception:
                continue

            if points:
                prev = points[-1]
                if (
                    abs(prev["lat"] - point["lat"]) < 1e-10
                    and
                    abs(prev["lon"] - point["lon"]) < 1e-10
                ):
                    continue

            points.append(point)

        if len(points) >= 2:
            routes[route_name] = points

    return routes


def sample_segment(
    a: dict,
    b: dict,
    step_nm: float,
):
    lon1, lat1 = float(a["lon"]), float(a["lat"])
    lon2, lat2 = float(b["lon"]), float(b["lat"])

    _, _, distance_m = GEOD.inv(
        lon1,
        lat1,
        lon2,
        lat2,
    )

    intervals = max(
        1,
        int(
            math.ceil(
                float(distance_m)
                / (max(0.5, step_nm) * 1852.0)
            )
        ),
    )

    points = [(lon1, lat1)]

    if intervals > 1:
        points.extend(
            (float(lon), float(lat))
            for lon, lat in GEOD.npts(
                lon1,
                lat1,
                lon2,
                lat2,
                intervals - 1,
            )
        )

    points.append((lon2, lat2))
    return points


def segment_intersects_area(
    a: dict,
    b: dict,
    area: dict,
    step_nm: float,
) -> bool:
    return any(
        point_is_inside_area(
            latitude=lat,
            longitude=lon,
            area=area,
        )
        for lon, lat in sample_segment(
            a,
            b,
            step_nm,
        )
    )


def route_parts_for_area(
    route_name: str,
    points: list[dict],
    area: dict,
    step_nm: float,
) -> list[dict]:
    """
    Return contiguous pieces whose route segments intersect the study area.
    Adjacent published route points are preserved so the line reaches the
    area boundary continuously.
    """
    parts = []
    current = []

    for a, b in zip(
        points[:-1],
        points[1:],
    ):
        if segment_intersects_area(
            a,
            b,
            area,
            step_nm,
        ):
            if not current:
                current = [a]

            if current[-1] != b:
                current.append(b)

        elif current:
            if len(current) >= 2:
                parts.append(current)
            current = []

    if current and len(current) >= 2:
        parts.append(current)

    return [
        {
            "name": route_name,
            "route": route_name,
            "type": "ATS",
            "segment_index": index,
            "points": part,
        }
        for index, part in enumerate(
            parts,
            start=1,
        )
    ]



# =============================================================================
# RAD APPENDIX 4 / DCT SUPPORT
# =============================================================================


def clean_text(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\n", " ").split()).strip()


def normalise_header(value) -> str:
    value = clean_text(value).upper()
    value = re.sub(r"[^A-Z0-9]+", " ", value)
    return " ".join(value.split())


def safe_float(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        try:
            out = float(value)
            return out if math.isfinite(out) else None
        except Exception:
            return None
    text_value = clean_text(value).upper().replace(",", ".")
    text_value = re.sub(r"^FL\s*", "", text_value)
    match = re.search(r"-?\d+(?:\.\d+)?", text_value)
    if not match:
        return None
    try:
        out = float(match.group(0))
        return out if math.isfinite(out) else None
    except Exception:
        return None


def point_code(value: str) -> str:
    value = clean_text(value).upper()
    value = re.sub(r"<\s*PT\s*>", " ", value)
    value = value.replace("/", " ")
    tokens = re.findall(r"[A-Z0-9]{2,8}", value)
    for token in tokens:
        if POINT_CODE_RE.fullmatch(token):
            return token
    return tokens[0] if tokens else ""


def _catalog_add(catalog, name, lat, lon, label=None, source=None):
    name = point_code(name)
    lat = safe_float(lat)
    lon = safe_float(lon)
    if not name or lat is None or lon is None:
        return
    item = {
        "name": name,
        "label": label or name,
        "lat": float(lat),
        "lon": float(lon),
    }
    if source:
        item["coordinate_source"] = source
    catalog.setdefault(name, item)


def load_point_catalogue(waypoints_path: Path, enr32_routes: dict) -> dict:
    catalog = {}
    for route_points in enr32_routes.values():
        for point in route_points:
            _catalog_add(
                catalog,
                point.get("name", ""),
                point.get("lat"),
                point.get("lon"),
                point.get("label"),
                "ENAIRE ENR 3.2",
            )

    if waypoints_path.exists():
        try:
            payload = json.loads(waypoints_path.read_text(encoding="utf-8"))
        except Exception:
            payload = {}
        if isinstance(payload, dict):
            for area_value in payload.values():
                if not isinstance(area_value, dict):
                    continue
                for point in area_value.get("waypoints", []):
                    if isinstance(point, dict):
                        _catalog_add(
                            catalog,
                            point.get("name", ""),
                            point.get("lat"),
                            point.get("lon"),
                            point.get("name"),
                            "waypoints.json",
                        )
    return catalog


def parse_airac_source_spec(value: str):
    if "=" not in value:
        raise argparse.ArgumentTypeError(
            "Expected AIRAC=PATH_OR_URL, e.g. 1901=/path/RAD_1901.xlsx"
        )
    cycle, source = value.split("=", 1)
    cycle, source = cycle.strip(), source.strip()
    if not re.fullmatch(r"\d{4}", cycle):
        raise argparse.ArgumentTypeError(f"Invalid AIRAC cycle: {cycle!r}")
    if not source:
        raise argparse.ArgumentTypeError("Empty RAD workbook source")
    return cycle, source


def read_workbook_bytes(source: str, timeout_s: float) -> bytes:
    if source.lower().startswith(("http://", "https://")):
        response = requests.get(
            source,
            headers={"User-Agent": "Mozilla/5.0 (OpenSky academic RAD downloader)"},
            timeout=timeout_s,
        )
        response.raise_for_status()
        return response.content
    path = Path(source).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"RAD workbook does not exist: {path}")
    if path.suffix.lower() == ".xls":
        raise RuntimeError(
            "Legacy .xls is not supported directly. Convert/save it as .xlsx first."
        )
    return path.read_bytes()


def detect_rad_header(row_values):
    headers = [normalise_header(value) for value in row_values]
    columns = {}
    for index, header in enumerate(headers):
        if not header:
            continue
        if header == "FROM" or header.startswith("FROM "):
            columns.setdefault("from", index)
        elif header == "TO" or header.startswith("TO "):
            columns.setdefault("to", index)
        elif "LOWER" in header and "VERTICAL" in header:
            columns.setdefault("lower_fl", index)
        elif "UPPER" in header and "VERTICAL" in header:
            columns.setdefault("upper_fl", index)
        elif "AVAILABLE" in header and "TIME" not in header:
            columns.setdefault("available", index)
        elif "UTILIZATION" in header:
            columns.setdefault("utilization", index)
        elif "TIME" in header and ("AVAILABILITY" in header or "APPLICABILITY" in header):
            columns.setdefault("time_availability", index)
        elif "ID" in header and ("NUMBER" in header or "NO" in header):
            columns.setdefault("id_number", index)
        elif "OPERATIONAL" in header and "GOAL" in header:
            columns.setdefault("operational_goal", index)
        elif "REMARK" in header:
            columns.setdefault("remark", index)
        elif "DIRECTION" in header and "CRUISING" in header:
            columns.setdefault("cruising_direction", index)
        elif "ATC" in header and "UNIT" in header:
            columns.setdefault("atc_unit", index)
        elif "CHANGE" in header and "RECORD" in header:
            columns.setdefault("change_record", index)
    return columns if "from" in columns and "to" in columns else None


def _cell(values, columns, key):
    index = columns.get(key)
    if index is None or index >= len(values):
        return None
    return values[index]


def parse_rad_workbook(workbook_bytes: bytes, airac: str, source_label: str) -> list[dict]:
    if load_workbook is None:
        raise RuntimeError(
            "openpyxl is required. Install with: python3 -m pip install openpyxl"
        )
    workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    records = []
    for sheet in workbook.worksheets:
        header_row = None
        columns = None
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 100), values_only=True),
            start=1,
        ):
            candidate = detect_rad_header(list(row))
            if candidate:
                header_row, columns = row_number, candidate
                break
        if header_row is None:
            continue

        blank_streak = 0
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            values = list(row)
            from_raw = clean_text(_cell(values, columns, "from"))
            to_raw = clean_text(_cell(values, columns, "to"))
            if not from_raw and not to_raw:
                blank_streak += 1
                if blank_streak >= 50:
                    break
                continue
            blank_streak = 0
            from_name, to_name = point_code(from_raw), point_code(to_raw)
            if not from_name or not to_name:
                continue
            records.append({
                "airac": airac,
                "source": source_label,
                "sheet": sheet.title,
                "from": from_name,
                "to": to_name,
                "lower_fl": safe_float(_cell(values, columns, "lower_fl")),
                "upper_fl": safe_float(_cell(values, columns, "upper_fl")),
                "available": clean_text(_cell(values, columns, "available")),
                "utilization": clean_text(_cell(values, columns, "utilization")),
                "time_availability": clean_text(_cell(values, columns, "time_availability")),
                "id_number": clean_text(_cell(values, columns, "id_number")),
                "operational_goal": clean_text(_cell(values, columns, "operational_goal")),
                "remark": clean_text(_cell(values, columns, "remark")),
                "cruising_direction": clean_text(_cell(values, columns, "cruising_direction")),
                "atc_unit": clean_text(_cell(values, columns, "atc_unit")),
                "change_record": clean_text(_cell(values, columns, "change_record")),
            })
    return records


def dct_is_drawable(record: dict, mode: str) -> bool:
    value = clean_text(record.get("available")).upper()
    yes = value in {"Y", "YES", "TRUE", "1", "AVAILABLE"}
    no = value in {"N", "NO", "FALSE", "0", "NOT AVAILABLE", "UNAVAILABLE"}
    if mode == "all":
        return True
    if mode == "not-no":
        return not no
    return yes


def resolve_dct_records(records, catalog, area, step_nm, availability_mode):
    resolved, unresolved = [], set()
    skipped_availability = 0
    for record in records:
        if not dct_is_drawable(record, availability_mode):
            skipped_availability += 1
            continue
        point_a = catalog.get(record["from"])
        point_b = catalog.get(record["to"])
        if point_a is None:
            unresolved.add(record["from"])
        if point_b is None:
            unresolved.add(record["to"])
        if point_a is None or point_b is None:
            continue
        if not segment_intersects_area(point_a, point_b, area, step_nm):
            continue
        item = dict(record)
        item["point_a"] = dict(point_a)
        item["point_b"] = dict(point_b)
        resolved.append(item)
    return resolved, unresolved, skipped_availability


def merge_dct_corridors(records):
    groups = defaultdict(list)
    for record in records:
        groups[tuple(sorted((record["from"], record["to"])))] .append(record)

    output = []
    for segment_index, (endpoints, group) in enumerate(sorted(groups.items()), start=1):
        name_a, name_b = endpoints
        first = group[0]
        by_name = {first["from"]: first["point_a"], first["to"]: first["point_b"]}
        cycles = sorted({record["airac"] for record in group})
        validity = []
        for cycle in cycles:
            if cycle in AIRAC_VALIDITY_2019:
                valid_from, valid_to = AIRAC_VALIDITY_2019[cycle]
                validity.append({"airac": cycle, "valid_from": valid_from, "valid_to": valid_to})

        conditions = []
        for record in sorted(group, key=lambda x: (x["airac"], x["from"], x["to"], x.get("id_number", ""))):
            condition = {
                k: v for k, v in record.items()
                if k not in {"point_a", "point_b", "source"} and v not in ("", None)
            }
            conditions.append(condition)

        output.append({
            "name": f"DCT {name_a}-{name_b}",
            "route": f"DCT {name_a}-{name_b}",
            "type": "DCT",
            "segment_index": segment_index,
            "points": [dict(by_name[name_a]), dict(by_name[name_b])],
            "airac_cycles": cycles,
            "validity": validity,
            "conditions": conditions,
        })
    return output


def load_existing(
    path: Path,
) -> dict:
    if not path.exists():
        return {}

    try:
        payload = json.loads(
            path.read_text(
                encoding="utf-8"
            )
        )
    except Exception:
        return {}

    return payload if isinstance(payload, dict) else {}


def save_output(
    path: Path,
    area_key: str,
    area: dict,
    ats_corridors: list[dict],
    dct_corridors: list[dict],
    enr32_source: dict | None,
    rad_sources: list[dict],
    unresolved_points: set[str],
    availability_mode: str,
):
    payload = load_existing(path)
    all_corridors = ats_corridors + dct_corridors
    cycles = sorted({c for corridor in dct_corridors for c in corridor.get("airac_cycles", [])})
    mixed = bool(cycles and enr32_source and not enr32_source.get("historical", False))

    payload["_meta"] = {
        "schema": "opensky-corridors-v2",
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "description": "ENAIRE ENR 3.2 ATS routes + EUROCONTROL RAD Appendix 4 DCTs.",
        "important": (
            "When several RAD AIRAC workbooks are supplied, corridors is the union "
            "of those cycles. DCT entries retain airac_cycles/validity metadata."
        ),
    }

    payload[area_key] = {
        "name": area.get("name", area_key),
        "sources": {"enr32": enr32_source, "rad_appendix4": rad_sources},
        "temporal": {
            "rad_airac_cycles": cycles,
            "mixed_temporal_sources": mixed,
            "warning": (
                "Historical RAD DCTs are mixed with a non-historical ENR 3.2 source."
                if mixed else None
            ),
        },
        "filter": {
            "area_key": area_key,
            "bbox": area["bbox"],
            "polygon": bool(area.get("polygon")),
            "rad_availability_mode": availability_mode,
            "description": (
                "ATS and DCT segments are retained when sampled geodesic points "
                "intersect the resolved study area."
            ),
        },
        "count": len(all_corridors),
        "counts": {"ATS": len(ats_corridors), "DCT": len(dct_corridors)},
        "unresolved_dct_points": sorted(unresolved_points),
        "corridors": all_corridors,
    }

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )


def parse_args():
    parser = argparse.ArgumentParser(
        description="Compile ENAIRE ATS routes plus EUROCONTROL RAD DCTs."
    )
    parser.add_argument("--area", default=DEFAULT_AREA)
    parser.add_argument("--source-url", default=DEFAULT_SOURCE_URL)
    parser.add_argument("--source-html", type=Path, default=None)
    parser.add_argument("--enr32-airac", default=None)
    parser.add_argument(
        "--rad-workbook",
        action="append",
        type=parse_airac_source_spec,
        default=[],
        metavar="AIRAC=PATH_OR_URL",
        help=(
            "RAD Appendix 4 workbook; repeat for several cycles. "
            "Example: --rad-workbook 1901=/path/RAD_1901.xlsx"
        ),
    )
    parser.add_argument(
        "--rad-availability",
        choices=("yes", "not-no", "all"),
        default="yes",
    )
    parser.add_argument("--waypoints", type=Path, default=DEFAULT_WAYPOINTS_FILE)
    parser.add_argument("--bbox-presets", type=Path, default=BBOX_PRESETS_FILE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_FILE)
    parser.add_argument("--area-sample-step-nm", type=float, default=5.0)
    parser.add_argument("--timeout", type=float, default=60.0)
    parser.add_argument("--skip-ats", action="store_true")
    parser.add_argument("--skip-dct", action="store_true")
    return parser.parse_args()


def main():
    args = parse_args()
    if args.area_sample_step_nm <= 0 or args.timeout <= 0:
        raise ValueError("Sampling step and timeout must be positive")

    presets = load_area_presets(args.bbox_presets)
    area = resolve_area_preset(args.area, presets)

    routes = {}
    ats_corridors = []
    enr32_source = None

    if not args.skip_ats:
        if args.source_html is not None:
            source_html = args.source_html.expanduser().resolve()
            if not source_html.exists():
                raise FileNotFoundError(f"Source HTML does not exist:\n{source_html}")
            html = source_html.read_text(encoding="utf-8", errors="replace")
            enr32_source = {
                "name": "Local ENR 3.2 HTML",
                "path": str(source_html),
                "historical": bool(args.enr32_airac),
                "airac": args.enr32_airac,
            }
        else:
            print("Downloading ENAIRE ENR 3.2:")
            print(args.source_url)
            html = fetch_html(args.source_url, args.timeout)
            enr32_source = {
                "name": "ENAIRE AIP ENR 3.2",
                "url": args.source_url,
                "historical": False,
                "airac": args.enr32_airac,
            }

        routes = parse_enr32_routes(html)
        for route_name, points in sorted(routes.items()):
            parts = route_parts_for_area(route_name, points, area, args.area_sample_step_nm)
            if args.enr32_airac:
                for part in parts:
                    part["airac_cycles"] = [args.enr32_airac]
            ats_corridors.extend(parts)

    print(f"Parsed ATS routes: {len(routes):,}")
    print(f"ATS corridor parts in {args.area}: {len(ats_corridors):,}")

    catalog = load_point_catalogue(args.waypoints.expanduser().resolve(), routes)
    print(f"Resolvable named points: {len(catalog):,}")

    all_rad_records = []
    rad_sources = []
    if not args.skip_dct:
        for cycle, source in args.rad_workbook:
            print()
            print(f"Loading RAD Appendix 4 AIRAC {cycle}:")
            print(f"  {source}")
            workbook_bytes = read_workbook_bytes(source, args.timeout)
            records = parse_rad_workbook(workbook_bytes, cycle, source)
            print(f"  Parsed point-to-point rows: {len(records):,}")
            all_rad_records.extend(records)
            meta = {"airac": cycle, "source": source, "record_count": len(records)}
            if cycle in AIRAC_VALIDITY_2019:
                meta["valid_from"], meta["valid_to"] = AIRAC_VALIDITY_2019[cycle]
            rad_sources.append(meta)

    if all_rad_records:
        resolved, unresolved, skipped = resolve_dct_records(
            all_rad_records,
            catalog,
            area,
            args.area_sample_step_nm,
            args.rad_availability,
        )
    else:
        resolved, unresolved, skipped = [], set(), 0

    dct_corridors = merge_dct_corridors(resolved)

    print()
    print(f"Drawable DCT records intersecting {args.area}: {len(resolved):,}")
    print(f"Unique DCT corridor lines: {len(dct_corridors):,}")
    print(f"RAD rows skipped by availability filter: {skipped:,}")
    print(f"Unresolved DCT endpoint names: {len(unresolved):,}")
    if unresolved:
        preview = ", ".join(sorted(unresolved)[:40])
        if len(unresolved) > 40:
            preview += ", ..."
        print(f"  {preview}")

    output = args.output.expanduser().resolve()
    save_output(
        output,
        args.area,
        area,
        ats_corridors,
        dct_corridors,
        enr32_source,
        rad_sources,
        unresolved,
        args.rad_availability,
    )

    print()
    print("=" * 78)
    print("CORRIDOR BUILD COMPLETE")
    print("=" * 78)
    print(f"Area            : {args.area}")
    print(f"ATS corridors   : {len(ats_corridors):,}")
    print(f"DCT corridors   : {len(dct_corridors):,}")
    print(f"Total corridors : {len(ats_corridors) + len(dct_corridors):,}")
    print(f"Output          : {output}")

    cycles = [cycle for cycle, _ in args.rad_workbook]
    if cycles:
        print(f"RAD AIRACs      : {', '.join(cycles)}")

    if args.area == "LECMSAN" and any(
        cycle in {"1813", "1901", "1902", "1903"} for cycle in cycles
    ):
        print()
        print("LECMSAN historical note:")
        print(
            "  For 2019-01-01..2019-02-28, use AIRAC 1813 for 1-2 Jan, "
            "then 1901, 1902 and 1903."
        )

    if cycles and enr32_source and not enr32_source.get("historical", False):
        print()
        print(
            "WARNING: historical RAD DCTs are being mixed with the current "
            "ENAIRE ENR 3.2. Use --source-html with historical ENR 3.2 when available."
        )


if __name__ == "__main__":
    main()
